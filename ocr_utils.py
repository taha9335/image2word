from PIL import Image
import pytesseract
import os
import cv2
import numpy as np
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import math

# Configure Tesseract path for Windows if not in PATH
if os.name == 'nt':
    tesseract_path = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    if os.path.exists(tesseract_path):
        pytesseract.pytesseract.tesseract_cmd = tesseract_path


def _pil_to_cv2(img: Image.Image):
    return cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)


def preprocess_image(pil_img: Image.Image):
    img = pil_img.convert('RGB')
    cv = _pil_to_cv2(img)
    gray = cv2.cvtColor(cv, cv2.COLOR_BGR2GRAY)
    gray = cv2.medianBlur(gray, 3)
    return gray


def _detect_bold(word_img):
    if word_img.size == 0:
        return False
    _, th = cv2.threshold(word_img, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    # foreground is dark text -> invert
    fg = (th == 0).astype(np.uint8)
    ratio = fg.sum() / (fg.shape[0] * fg.shape[1])
    # rough heuristic: thicker strokes produce higher foreground density
    return ratio > 0.20


def _detect_italic(word_img):
    coords = cv2.findNonZero(255 - word_img)  # nonzero points of text
    if coords is None or len(coords) < 5:
        return False
    coords = coords.reshape(-1, 2).astype(np.float32)
    mean, eigenvectors = cv2.PCACompute(coords, mean=None, maxComponents=2)
    vec = eigenvectors[0]
    angle = math.degrees(math.atan2(vec[1], vec[0]))
    return abs(angle) > 8


def _line_alignment(line_words, image_width):
    """Determine line alignment - defaults to LEFT for cleaner output."""
    if not line_words:
        return WD_PARAGRAPH_ALIGNMENT.LEFT
    
    lefts = [w['left'] for w in line_words]
    rights = [w['left'] + w['width'] for w in line_words]
    left = min(lefts)
    right = max(rights)
    text_width = right - left
    
    # Only detect center if text is short (like a heading) and truly centered
    left_margin = left
    right_margin = image_width - right
    
    # Text must be less than 60% of image width and margins roughly equal for center
    is_short_text = text_width < image_width * 0.6
    margins_balanced = abs(left_margin - right_margin) < image_width * 0.1
    
    if is_short_text and margins_balanced and left_margin > image_width * 0.15:
        return WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Right alignment: large left margin, small right margin
    if left_margin > image_width * 0.4 and right_margin < image_width * 0.1:
        return WD_PARAGRAPH_ALIGNMENT.RIGHT
    
    # Default to left alignment for normal text
    return WD_PARAGRAPH_ALIGNMENT.LEFT


def process_image_to_docx(input_path, output_path):
    # Allow overriding system tesseract path with env var `TESSERACT_CMD`
    tpath = os.environ.get('TESSERACT_CMD') or os.environ.get('TESSERACT_PATH')
    if tpath:
        pytesseract.pytesseract.tesseract_cmd = tpath

    pil = Image.open(input_path).convert('RGB')
    gray = preprocess_image(pil)
    h, w = gray.shape

    try:
        data = pytesseract.image_to_data(pil, output_type=pytesseract.Output.DICT)
    except pytesseract.pytesseract.TesseractNotFoundError:
        raise RuntimeError("tesseract is not installed or it's not in your PATH. See README file for more information.")

    # Group words by (block_num, par_num, line_num)
    lines = {}
    n = len(data['text'])
    for i in range(n):
        txt = data['text'][i].strip()
        if txt == '':
            continue
        key = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        entry = {
            'text': txt,
            'left': int(data['left'][i]),
            'top': int(data['top'][i]),
            'width': int(data['width'][i]),
            'height': int(data['height'][i]),
            'conf': int(data['conf'][i]) if str(data['conf'][i]).lstrip('-').isdigit() else -1,
        }
        lines.setdefault(key, []).append(entry)

    # Convert to ordered list of lines
    ordered_lines = []
    for key, words in lines.items():
        top = min(w['top'] for w in words)
        ordered_lines.append((top, words))
    ordered_lines.sort(key=lambda x: x[0])

    # Group lines into paragraphs by vertical gaps
    paragraphs = []
    if ordered_lines:
        cur_para = [ordered_lines[0][1]]
        last_top = ordered_lines[0][0]
        heights = [max(w['height'] for w in ln) for _, ln in ordered_lines]
        median_h = np.median(heights) if heights else 0
        for top, words in ordered_lines[1:]:
            gap = top - last_top
            if median_h and gap > median_h * 1.8:
                paragraphs.append(cur_para)
                cur_para = [words]
            else:
                cur_para.append(words)
            last_top = top
        if cur_para:
            paragraphs.append(cur_para)

    doc = Document()

    for para in paragraphs:
        # determine paragraph alignment by majority of line alignments
        align_counts = {}
        line_aligns = []
        for line in para:
            align = _line_alignment(line, w)
            line_aligns.append(align)
            align_counts[align] = align_counts.get(align, 0) + 1
        # pick majority
        para_align = max(align_counts.items(), key=lambda x: x[1])[0]

        p = doc.add_paragraph()
        p.alignment = para_align

        for li, line in enumerate(para):
            # sort words by left
            line = sorted(line, key=lambda x: x['left'])
            for wi, word in enumerate(line):
                # crop word image to detect bold/italic
                l, t, ww, hh = word['left'], word['top'], word['width'], word['height']
                # clamp
                l0 = max(0, l)
                t0 = max(0, t)
                r0 = min(w, l + ww)
                b0 = min(h, t + hh)
                crop = gray[t0:b0, l0:r0]
                if crop.size == 0:
                    is_bold = False
                    is_italic = False
                else:
                    is_bold = _detect_bold(crop)
                    is_italic = _detect_italic(crop)

                run = p.add_run(word['text'])
                run.bold = bool(is_bold)
                run.italic = bool(is_italic)
                if wi < len(line) - 1:
                    run.add_text(' ')
            if li < len(para) - 1:
                run = p.add_run()
                run.add_break()

    doc.save(output_path)
